import os
import re
from datetime import datetime
from werkzeug.utils import secure_filename
import pdfplumber
from docx import Document
import openpyxl

class FileHandler:
    """مدیریت آپلود و استخراج متن از فایل‌ها"""
    
    ALLOWED_EXTENSIONS = {'txt', 'rtf', 'pdf', 'docx', 'xlsx', 'xls'}
    
    def __init__(self, upload_folder):
        self.upload_folder = upload_folder
        os.makedirs(upload_folder, exist_ok=True)
    
    def allowed_file(self, filename):
        """بررسی پسوند مجاز فایل"""
        return '.' in filename and \
               filename.rsplit('.', 1)[1].lower() in self.ALLOWED_EXTENSIONS
    
    def save_file(self, file):
        """ذخیره فایل آپلود شده با نام یکتا"""
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        original_name = secure_filename(file.filename)
        safe_filename = f"{timestamp}_{original_name}"
        file_path = os.path.join(self.upload_folder, safe_filename)
        
        file.save(file_path)
        
        return {
            'name': original_name,
            'path': file_path,
            'size': os.path.getsize(file_path),
            'type': original_name.rsplit('.', 1)[1].lower() if '.' in original_name else 'unknown',
            'saved_as': safe_filename
        }
    
    def extract_text(self, file_path, original_filename=None):
        """استخراج متن از فایل با فرمت‌های مختلف"""
        filename = original_filename or os.path.basename(file_path)
        ext = filename.lower().split('.')[-1] if '.' in filename else ''
        
        try:
            with open(file_path, 'rb') as f:
                content = f.read()
            
            if ext == 'rtf':
                return self._extract_from_rtf(content)
            elif ext == 'txt':
                return self._extract_from_txt(content)
            elif ext == 'pdf':
                return self._extract_from_pdf(file_path)
            elif ext == 'docx':
                return self._extract_from_docx(file_path)
            elif ext in ['xlsx', 'xls']:
                return self._extract_from_excel(file_path)
            else:
                return self._extract_generic(content)
                
        except Exception as e:
            raise Exception(f"خطا در استخراج متن از فایل: {str(e)}")
    
    def _extract_from_rtf(self, content):
        """استخراج از RTF"""
        try:
            # تلاش با encoding‌های مختلف
            for encoding in ['utf-8', 'windows-1256', 'latin-1']:
                try:
                    text = content.decode(encoding, errors='ignore')
                    break
                except:
                    continue
            else:
                text = content.decode('utf-8', errors='ignore')
            
            # پاکسازی دستورات RTF
            text = re.sub(r'\\[a-z]+\d*[\s\-]?', ' ', text)
            text = re.sub(r'\{|\}', ' ', text)
            text = re.sub(r'\s+', ' ', text)
            
            return text.strip()
            
        except Exception as e:
            raise Exception(f"خطا در پردازش RTF: {e}")
    
    def _extract_from_txt(self, content):
        """استخراج از TXT با encoding‌های مختلف"""
        encodings = ['utf-8', 'cp1256', 'windows-1256', 'iso-8859-1']
        
        for encoding in encodings:
            try:
                return content.decode(encoding).strip()
            except:
                continue
        
        return content.decode('utf-8', errors='ignore').strip()
    
    def _extract_from_pdf(self, file_path):
        """استخراج از PDF"""
        text = ""
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        return text.strip()
    
    def _extract_from_docx(self, file_path):
        """استخراج از DOCX"""
        doc = Document(file_path)
        return "\n".join([para.text for para in doc.paragraphs]).strip()
    
    def _extract_from_excel(self, file_path):
        """استخراج از Excel"""
        wb = openpyxl.load_workbook(file_path)
        text = ""
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = " ".join([str(cell) for cell in row if cell])
                if row_text:
                    text += row_text + "\n"
        return text.strip()
    
    def _extract_generic(self, content):
        """استخراج عمومی برای فرمت‌های ناشناخته"""
        encodings = ['utf-8', 'cp1256', 'windows-1256', 'latin-1']
        
        for encoding in encodings:
            try:
                return content.decode(encoding).strip()
            except:
                continue
        
        return content.decode('utf-8', errors='ignore').strip()
    
    def delete_file(self, file_path):
        """حذف فایل"""
        try:
            if os.path.exists(file_path):
                os.remove(file_path)
                return True
        except Exception as e:
            print(f"خطا در حذف فایل: {e}")
        return False